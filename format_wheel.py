"""
Format Wheel v4 - GTA-style radial file converter (Windows)

Run:      python format_wheel.py
Required: pip install PyQt6 pynput Pillow python-docx openpyxl pywin32
Optional (better results):
    pip install pymupdf        # real PDF page images, better PDF text + tables
    pip install pdf2docx       # layout-preserving PDF -> DOCX
    pip install reportlab      # real-text PDFs (selectable) for text/Office -> PDF
    pip install pypdf xlrd     # extra fallbacks for PDF text / legacy .xls
    pip install pytesseract    # OCR (also install Tesseract) for image/scanned PDF -> TXT/CSV
MS Word / Excel (via pywin32) or LibreOffice give pixel-perfect DOCX/XLSX -> PDF/images
and let the app read legacy .doc / .xls files.
"""
from __future__ import annotations

import sys, os, re, io, csv, json, math, shutil, tempfile, textwrap, itertools
import subprocess, zipfile, traceback, codecs, functools, importlib.util
from pathlib import Path
from datetime import datetime, date
from html.parser import HTMLParser

# ------------------------------------------------------------------ Qt
try:
    from PyQt6.QtWidgets import (
        QApplication, QWidget, QSystemTrayIcon, QMenu, QFileDialog, QDialog,
        QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QCheckBox, QLineEdit,
        QColorDialog,
    )
    from PyQt6.QtCore import (
        Qt, QPoint, QPointF, QRect, QRectF, QTimer, pyqtSignal, QMimeData,
        QUrl, QObject, QThread, QLockFile,
    )
    from PyQt6.QtGui import (
        QPainter, QColor, QPen, QBrush, QFont, QIcon, QPixmap, QAction,
        QRadialGradient, QLinearGradient, QCursor, QPainterPath, QFontMetrics,
    )
except ImportError:
    print("PyQt6 required: pip install PyQt6")
    sys.exit(1)

# ------------------------------------------------------------------ optional libs
try:
    from PIL import Image, ImageDraw, ImageFont, ImageOps, ImageSequence
    PIL_OK = True
except ImportError:
    PIL_OK = False

try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

try:
    import pythoncom
    import win32com.client
    WIN32_OK = True
except ImportError:
    WIN32_OK = False

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    csv.field_size_limit(2 ** 31 - 1)
except Exception:
    pass

# ------------------------------------------------------------------ config
CONFIG_DIR = Path(os.getenv("APPDATA") or Path.home()) / "FormatWheel"
CONFIG_DIR.mkdir(parents=True, exist_ok=True)
CONFIG_FILE = CONFIG_DIR / "config.json"
LOG_FILE = CONFIG_DIR / "format_wheel.log"

DEFAULT_CONFIG = {
    "hotkey": "ctrl+shift+c",
    "output_dir": str(Path.home() / "Documents" / "FormatWheel"),
    "theme_color": "#00e5ff",
    "start_with_windows": False,
    "copy_after_convert": False,
    "auto_convert_clipboard": True,
    "wheel_radius": 250,        # outer radius of the wheel in px
    "backdrop_dim": 150,        # 40-230, darkness of the screen overlay
    "release_to_select": True,  # GTA style: hold hotkey, aim, release
    "open_folder_after": False,
    "quiet_success": False,
}


def _log(msg):
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}\n")
    except Exception:
        pass


def load_config():
    if CONFIG_FILE.exists():
        try:
            return {**DEFAULT_CONFIG, **json.loads(CONFIG_FILE.read_text(encoding="utf-8"))}
        except Exception:
            pass
    return DEFAULT_CONFIG.copy()


def save_config(cfg):
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2), encoding="utf-8")


CONFIG = load_config()
Path(CONFIG["output_dir"]).mkdir(parents=True, exist_ok=True)

# ------------------------------------------------------------------ format registry
FORMATS = {
    "PDF":  {"ext": "pdf",  "color": "#ff3b30", "desc": "Portable Document"},
    "DOCX": {"ext": "docx", "color": "#2f7bff", "desc": "Word Document"},
    "XLSX": {"ext": "xlsx", "color": "#30b14b", "desc": "Excel Spreadsheet"},
    "CSV":  {"ext": "csv",  "color": "#ff2d55", "desc": "Comma-Separated Values"},
    "TXT":  {"ext": "txt",  "color": "#c19a6b", "desc": "Plain Text"},
    "ZIP":  {"ext": "zip",  "color": "#7a78ff", "desc": "Compressed Archive"},
    "TIFF": {"ext": "tiff", "color": "#ffcc00", "desc": "TIFF (multi-page)"},
    "BMP":  {"ext": "bmp",  "color": "#9aa0ad", "desc": "Bitmap Image"},
    "WEBP": {"ext": "webp", "color": "#5ac8fa", "desc": "WebP Image"},
    "GIF":  {"ext": "gif",  "color": "#bf5af2", "desc": "GIF Image"},
    "ICO":  {"ext": "ico",  "color": "#ff6ac1", "desc": "Windows Icon"},
    "PNG":  {"ext": "png",  "color": "#34c759", "desc": "Lossless Image"},
    "JPG":  {"ext": "jpg",  "color": "#ff9500", "desc": "JPEG Image"},
}
WHEEL_ORDER = list(FORMATS.keys())  # clockwise from the top
IMG_TARGETS = {"JPG", "PNG", "GIF", "WEBP", "BMP", "TIFF", "ICO"}
PAGE_TARGETS = IMG_TARGETS | {"PDF"}
PIL_FMT = {"JPG": "JPEG", "PNG": "PNG", "GIF": "GIF", "WEBP": "WEBP",
           "BMP": "BMP", "TIFF": "TIFF", "PDF": "PDF", "ICO": "ICO"}

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff", ".tif", ".ico"}
TEXT_EXTS = {
    ".txt", ".md", ".log", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".rst",
    ".yaml", ".yml", ".ini", ".cfg", ".conf", ".toml", ".env",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs",
    ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".css", ".scss",
    ".less", ".vue", ".swift", ".kt", ".r", ".m", ".lua", ".pl", ".dart", ".scala",
    ".asm", ".vbs", ".make", ".cmake", ".gitignore", ".dockerfile", ".editorconfig",
}
PROSE_EXTS = {".txt", ".md", ".log", ".rst"}
DOC_EXTS = {".docx", ".doc"}
XLS_EXTS = {".xlsx", ".xls", ".xlsm"}
PDF_EXTS = {".pdf"}

MAX_PAGES = 200
MAX_XLSX_ROWS = 1_048_576
MONO_FONTS = ("consola.ttf", "DejaVuSansMono.ttf", "LiberationMono-Regular.ttf",
              "cour.ttf", "Menlo.ttc")


# ================================================================== generic helpers
def _unique_path(out_dir: Path, base: str, ext: str) -> Path:
    base = re.sub(r'[\\/:*?"<>|]+', "_", base).strip() or "file"
    out = out_dir / f"{base}.{ext}"
    i = 1
    while out.exists():
        out = out_dir / f"{base}_{i}.{ext}"
        i += 1
    return out


_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _clean(s) -> str:
    return _XML_BAD.sub("", str(s))


_NUM_RE = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d+)?$")


def _auto_num(v):
    if isinstance(v, str):
        s = v.strip()
        if _NUM_RE.match(s):
            return float(s) if "." in s else int(s)
    return v


def _norm_ext(ext: str) -> str:
    e = ext.lower().lstrip(".")
    return {"jpeg": "jpg", "tif": "tiff"}.get(e, e)


def _read_text_any(src: Path) -> str:
    """Decode text files without ever crashing or producing mojibake."""
    raw = src.read_bytes()
    if raw.startswith(codecs.BOM_UTF8):
        return raw[3:].decode("utf-8", errors="replace")
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _looks_like_text(head: bytes) -> bool:
    if not head:
        return True
    if head.startswith((codecs.BOM_UTF8, codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return True
    if b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as e:
        if e.start >= len(head) - 4:
            return True
    good = sum(1 for b in head if 32 <= b < 127 or b in (9, 10, 13) or b >= 128)
    return good / len(head) > 0.9


@functools.lru_cache(maxsize=256)
def _source_kind(path: str) -> str:
    """'image' | 'text' | 'docx' | 'xlsx' | 'pdf' | 'folder' | 'unknown'."""
    src = Path(path)
    if src.is_dir():
        return "folder"
    ext = src.suffix.lower()
    if ext in IMAGE_EXTS:
        return "image"
    if ext in TEXT_EXTS:
        return "text"
    if ext in DOC_EXTS:
        return "docx"
    if ext in XLS_EXTS:
        return "xlsx"
    if ext in PDF_EXTS:
        return "pdf"
    try:
        with open(src, "rb") as f:
            head = f.read(4096)
    except Exception:
        return "unknown"
    if head.startswith(b"%PDF"):
        return "pdf"
    if PIL_OK:
        try:
            with Image.open(src) as im:
                im.verify()
            return "image"
        except Exception:
            pass
    return "text" if _looks_like_text(head) else "unknown"


_FONT_DIRS = [
    Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts",
    Path("/usr/share/fonts"), Path("/Library/Fonts"), Path("/System/Library/Fonts"),
    Path.home() / "Library/Fonts", Path.home() / ".fonts",
]


@functools.lru_cache(maxsize=8)
def _find_font(names: tuple):
    for n in names:
        for d in _FONT_DIRS:
            if not d.exists():
                continue
            p = d / n
            if p.exists():
                return str(p)
            try:
                hit = next(d.rglob(n), None)
                if hit:
                    return str(hit)
            except Exception:
                pass
    return None


@functools.lru_cache(maxsize=1)
def _find_soffice():
    p = shutil.which("soffice") or shutil.which("soffice.exe")
    if p:
        return p
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/usr/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(c):
            return c
    return None


def _has_reportlab():
    return importlib.util.find_spec("reportlab") is not None


_OCR_STATE = None


def ocr_available():
    global _OCR_STATE
    if _OCR_STATE is None:
        try:
            import pytesseract
            if not shutil.which("tesseract"):
                for p in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                          r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
                    if os.path.exists(p):
                        pytesseract.pytesseract.tesseract_cmd = p
                        break
            pytesseract.get_tesseract_version()
            _OCR_STATE = True
        except Exception:
            _OCR_STATE = False
    return _OCR_STATE


def _ocr(img) -> str:
    import pytesseract
    return pytesseract.image_to_string(img)


# ================================================================== document model
class Doc:
    """Format-neutral intermediate representation: headings, paragraphs, tables, images."""

    def __init__(self, mono=False):
        self.blocks = []
        self.mono = mono

    def h(self, text, level=1): self.blocks.append(("h", text, level))
    def p(self, text): self.blocks.append(("p", text))
    def t(self, rows, name="Table"): self.blocks.append(("t", name, rows))
    def img(self, path): self.blocks.append(("img", str(path)))
    def pb(self): self.blocks.append(("pb",))

    @property
    def tables(self):
        return [b for b in self.blocks if b[0] == "t" and b[2]]

    @property
    def text_lines(self):
        return [l for b in self.blocks if b[0] in ("h", "p") for l in b[1].split("\n")]

    @property
    def has_text(self):
        return any(l.strip() for l in self.text_lines)

    def lines(self, aligned=False):
        out = []
        for b in self.blocks:
            k = b[0]
            if k == "h":
                out += ["", b[1], ""]
            elif k == "p":
                out.extend(b[1].split("\n"))
            elif k == "t":
                out += _aligned(b[2]) if aligned else ["\t".join(map(str, r)) for r in b[2]]
                out.append("")
            elif k == "pb":
                out.append("")
        return out


def _aligned(rows):
    if not rows:
        return []
    n = max(len(r) for r in rows)
    cells = [[("" if c is None else str(c)).replace("\n", " ") for c in r] + [""] * (n - len(r))
             for r in rows]
    w = [min(28, max(len(r[i]) for r in cells)) for i in range(n)]
    lines = [" | ".join(c[:w[i]].ljust(w[i]) for i, c in enumerate(r)).rstrip() for r in cells]
    if len(lines) > 1:
        lines.insert(1, "-+-".join("-" * x for x in w))
    return lines


def _fmt_cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) \
            else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ================================================================== readers -> Doc
class _TextExtractor(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
             "section", "article", "table", "ul", "ol", "pre"}
    SKIP = {"script", "style", "head", "noscript", "template"}

    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, d):
        if not self.skip:
            self.parts.append(d)

    def text(self):
        t = "".join(self.parts)
        t = re.sub(r"[ \t]+", " ", t)
        return re.sub(r"\n\s*\n\s*\n+", "\n\n", t).strip()


def _parse_delimited(text, ext):
    delim = "\t" if ext == ".tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _read_text_doc(src: Path, ext: str) -> Doc:
    text = _read_text_any(src)
    if ext in (".csv", ".tsv"):
        d = Doc()
        d.t(_parse_delimited(text, ext), src.stem)
        return d
    if ext in (".html", ".htm"):
        ex = _TextExtractor()
        ex.feed(text)
        d = Doc()
        for l in ex.text().split("\n"):
            d.p(l)
        return d
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    d = Doc(mono=ext not in PROSE_EXTS)
    for l in lines:
        m = re.match(r"^(#{1,6})\s+(.*\S)\s*$", l) if ext == ".md" else None
        if m:
            d.h(m.group(2), len(m.group(1)))
        else:
            d.p(l)
    return d


def _read_docx_doc(src: Path, notes) -> Doc:
    if not DOCX_OK:
        raise RuntimeError("python-docx not installed (pip install python-docx)")
    d = docx.Document(str(src))
    doc = Doc()
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = DocxParagraph(child, d)
            try:
                sname = para.style.name if para.style is not None else ""
            except Exception:
                sname = ""
            if sname == "Title":
                doc.h(para.text, 0)
            elif sname.startswith("Heading"):
                try:
                    lvl = int(sname.split()[-1])
                except ValueError:
                    lvl = 1
                doc.h(para.text, lvl)
            else:
                doc.p(para.text)
        elif tag == "tbl":
            tbl = DocxTable(child, d)
            doc.t([[c.text for c in row.cells] for row in tbl.rows], f"Table {len(doc.tables) + 1}")
    try:
        if len(d.inline_shapes):
            notes.append("embedded images were not carried over")
    except Exception:
        pass
    return doc


def _read_xlsx_doc(src: Path, notes) -> Doc:
    doc = Doc()
    if src.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise RuntimeError("legacy .xls needs MS Excel, LibreOffice or xlrd (pip install xlrd)")
        book = xlrd.open_workbook(str(src))
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                        except Exception:
                            pass
                    row.append(_fmt_cell(v))
                rows.append(row)
            if rows:
                doc.t(rows, sh.name)
        return doc
    if not XLSX_OK:
        raise RuntimeError("openpyxl not installed (pip install openpyxl)")
    wb = openpyxl.load_workbook(str(src), data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            rows = [[_fmt_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            while rows and not any(rows[-1]):
                rows.pop()
            ncols = max((max((i for i, c in enumerate(r) if c != ""), default=-1) for r in rows),
                        default=-1) + 1
            rows = [r[:ncols] for r in rows]
            if rows and ncols:
                doc.t(rows, ws.title)
    finally:
        wb.close()
    if not doc.tables:
        notes.append("workbook contains no data")
    return doc


def _fitz_pages(src, notes, zoom=2.0):
    d = fitz.open(str(src))
    try:
        if d.needs_pass and not d.authenticate(""):
            raise RuntimeError("PDF is password-protected")
        for i, page in enumerate(d):
            if i >= MAX_PAGES:
                notes.append(f"only the first {MAX_PAGES} pages were converted")
                break
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            yield Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    finally:
        d.close()


def _pdf_page_images(pdf, notes, zoom=2.0):
    """Iterator of PIL page images, or None when no rasteriser is available."""
    if not PIL_OK:
        return None
    if fitz is not None:
        return _fitz_pages(pdf, notes, zoom)
    try:
        from pdf2image import convert_from_path
        return iter(convert_from_path(str(pdf), dpi=int(72 * zoom), last_page=MAX_PAGES))
    except Exception:
        return None


def _read_pdf_doc(src: Path, notes, want_tables, embed_if_empty, tmp: Path) -> Doc:
    doc = Doc()
    if fitz is not None:
        d = fitz.open(str(src))
        try:
            if d.needs_pass and not d.authenticate(""):
                raise RuntimeError("PDF is password-protected")
            use_ocr = None
            total = len(d)
            for i, page in enumerate(d):
                try:
                    text = page.get_text("text", sort=True)
                except TypeError:
                    text = page.get_text("text")
                if not text.strip() and PIL_OK:
                    if use_ocr is None:
                        use_ocr = ocr_available()
                    if use_ocr:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        text = _ocr(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
                for para in re.split(r"\n\s*\n", text.strip()):
                    if para.strip():
                        doc.p(para.strip())
                if want_tables:
                    try:
                        for n, tb in enumerate(page.find_tables().tables, 1):
                            rows = [[("" if c is None else str(c)) for c in r] for r in tb.extract()]
                            if rows:
                                doc.t(rows, f"p{i + 1}_table{n}")
                    except Exception:
                        pass
                if i < total - 1:
                    doc.pb()
        finally:
            d.close()
    elif pypdf is not None:
        r = pypdf.PdfReader(str(src))
        if r.is_encrypted and not r.decrypt(""):
            raise RuntimeError("PDF is password-protected")
        for i, pg in enumerate(r.pages):
            for para in re.split(r"\n\s*\n", (pg.extract_text() or "").strip()):
                if para.strip():
                    doc.p(para.strip())
            if i < len(r.pages) - 1:
                doc.pb()
    else:
        raise RuntimeError("PDF support needs PyMuPDF or pypdf (pip install pymupdf)")

    if not doc.has_text and not doc.tables:
        if embed_if_empty and fitz is not None and PIL_OK:
            doc = Doc()
            for n, img in enumerate(_fitz_pages(src, notes, 1.5)):
                pth = tmp / f"pdfpage_{n:04d}.png"
                img.save(pth)
                doc.img(pth)
            notes.append("scanned PDF - pages were inserted as images")
        else:
            notes.append("PDF has no extractable text (scanned?) - install pytesseract + Tesseract for OCR")
    return doc


def _image_doc(src: Path, target: str, tmp: Path) -> Doc:
    if not PIL_OK:
        raise RuntimeError("Pillow not installed (pip install Pillow)")
    doc = Doc()
    with Image.open(src) as im:
        base = _prep(ImageOps.exif_transpose(im), "PNG")
        if target in ("DOCX", "XLSX"):
            pth = tmp / "image.png"
            base.save(pth)
            doc.img(pth)
        else:
            if not ocr_available():
                raise RuntimeError("image -> text needs OCR (pip install pytesseract + install Tesseract)")
            for l in _ocr(_flatten(base)).splitlines():
                doc.p(l)
    return doc


# ================================================================== Office / LibreOffice bridges
def _com_convert(src: Path, fmt: str, tmp: Path):
    if not WIN32_OK:
        return None
    ext = src.suffix.lower()
    out = tmp / f"{src.stem}.{fmt}"
    app = None
    try:
        if ext in DOC_EXTS:
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            d = app.Documents.Open(str(src.resolve()), ReadOnly=True,
                                   AddToRecentFiles=False, ConfirmConversions=False)
            try:
                d.SaveAs2(str(out), FileFormat=17 if fmt == "pdf" else 16)
            finally:
                d.Close(False)
        elif ext in XLS_EXTS:
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            wb = app.Workbooks.Open(str(src.resolve()), ReadOnly=True)
            try:
                if fmt == "pdf":
                    wb.ExportAsFixedFormat(0, str(out))
                else:
                    wb.SaveAs(str(out), FileFormat=51)
            finally:
                wb.Close(False)
        else:
            return None
        return out if out.exists() and out.stat().st_size > 0 else None
    except Exception as e:
        _log(f"COM convert failed: {e}")
        return None
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass


def _lo_convert(src: Path, fmt: str, tmp: Path):
    exe = _find_soffice()
    if not exe:
        return None
    try:
        outdir = tmp / f"lo_{fmt}"
        outdir.mkdir(exist_ok=True)
        profile = (tmp / "lo_profile").resolve().as_uri()
        cmd = [exe, f"-env:UserInstallation={profile}", "--headless",
               "--convert-to", fmt, "--outdir", str(outdir), str(src)]
        subprocess.run(cmd, capture_output=True, timeout=180,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        res = outdir / f"{src.stem}.{fmt}"
        return res if res.exists() else None
    except Exception as e:
        _log(f"LibreOffice convert failed: {e}")
        return None


def _office_convert(src: Path, fmt: str, tmp: Path):
    return _com_convert(src, fmt, tmp) or _lo_convert(src, fmt, tmp)


# ================================================================== image helpers
def _flatten(img, bg=(255, 255, 255)):
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        base = Image.new("RGB", rgba.size, bg)
        base.paste(rgba, mask=rgba.split()[3])
        return base
    return img.convert("RGB")


def _prep(img, target):
    if target in ("JPG", "BMP", "PDF"):
        return _flatten(img)
    if img.mode not in ("1", "L", "LA", "P", "RGB", "RGBA"):
        has_a = "A" in img.mode or "transparency" in img.info
        return img.convert("RGBA" if has_a else "RGB")
    return img


def _save_frames(frames, out: Path, target, duration=100):
    first, rest = frames[0], frames[1:]
    kw = {}
    if target == "JPG":
        kw = dict(quality=95, optimize=True, subsampling=0)
    elif target == "WEBP":
        kw = dict(quality=95, method=4)
    elif target == "TIFF":
        kw = dict(compression="tiff_lzw")
    elif target == "PDF":
        kw = dict(resolution=max(72.0, max(first.size) / 11.69))
    elif target == "GIF":
        kw = dict(loop=0, duration=duration, disposal=2)
    elif target == "ICO":
        im = first.convert("RGBA")
        side = max(im.size)
        sq = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        sq.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        first = sq.resize((256, 256), Image.LANCZOS)  # all icon sizes are generated from 256px
        kw = dict(sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
    if rest and target in ("GIF", "WEBP", "TIFF", "PDF"):
        kw.update(save_all=True, append_images=rest)
        if target == "WEBP":
            kw["duration"] = duration
    try:
        first.save(out, PIL_FMT[target], **kw)
    except (OSError, TypeError, ValueError):
        if target == "TIFF":
            kw.pop("compression", None)
            first.save(out, PIL_FMT[target], **kw)
        else:
            raise


def _save_pages(pages, out_dir: Path, stem: str, target: str, notes):
    """Save an iterator of PIL pages: one file for PDF/TIFF, one file per page otherwise."""
    it = iter(pages)
    first = next(it, None)
    if first is None:
        raise RuntimeError("nothing to render")
    second = next(it, None)
    ext = FORMATS[target]["ext"]

    def prep(i): return _prep(i, target)

    if second is None:
        out = _unique_path(out_dir, stem, ext)
        _save_frames([prep(first)], out, target)
        return [out]
    if target in ("PDF", "TIFF"):
        frames = [prep(first), prep(second)] + [prep(p) for p in it]
        out = _unique_path(out_dir, stem, ext)
        _save_frames(frames, out, target)
        return [out]
    base, k = stem, 1
    while (out_dir / f"{base}_p001.{ext}").exists():
        base, k = f"{stem}_{k}", k + 1
    outs = []
    for n, pg in enumerate(itertools.chain([first, second], it), 1):
        out = out_dir / f"{base}_p{n:03d}.{ext}"
        _save_frames([prep(pg)], out, target)
        outs.append(out)
    notes.append(f"{len(outs)} pages saved as separate {ext.upper()} files")
    return outs


def _convert_image(src: Path, target, out_dir: Path, notes):
    if not PIL_OK:
        raise RuntimeError("Pillow not installed (pip install Pillow)")
    frames, duration = [], 100
    with Image.open(src) as im:
        n = getattr(im, "n_frames", 1)
        duration = im.info.get("duration", 100) or 100
        keep_all = n > 1 and target in ("GIF", "WEBP", "TIFF", "PDF")
        for i, fr in enumerate(ImageSequence.Iterator(im)):
            if (not keep_all and i > 0) or i >= 500:
                break
            frames.append(_prep(ImageOps.exif_transpose(fr), target))
        if n > 1 and not keep_all:
            notes.append("animated/multi-page image - only the first frame was kept")
    out = _unique_path(out_dir, src.stem, FORMATS[target]["ext"])
    _save_frames(frames, out, target, duration)
    return [out]


# ================================================================== text rendering (fallback for PDF/images)
PAGE_W, PAGE_H, MARGIN = 1240, 1754, 90


def _wrap(line, cols):
    line = _clean(line).expandtabs(4)
    if len(line) <= cols:
        return [line]
    indent = " " * min(len(line) - len(line.lstrip(" ")), cols // 2)
    return textwrap.wrap(line, cols, replace_whitespace=False, subsequent_indent=indent,
                         break_long_words=True) or [""]


def _load_font(size):
    path = _find_font(MONO_FONTS)
    if path:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            pass
    try:
        return ImageFont.load_default(size)
    except TypeError:
        return ImageFont.load_default()


def _render_pages(lines, notes):
    size = 20
    font, small = _load_font(size), _load_font(14)
    try:
        cw = font.getlength("M")
    except Exception:
        cw = size * 0.6
    lh = int(size * 1.5)
    cols = max(20, int((PAGE_W - 2 * MARGIN) / max(cw, 1)))
    rows = (PAGE_H - 2 * MARGIN) // lh
    wrapped = []
    for ln in lines:
        wrapped.extend(_wrap(ln, cols))
    while wrapped and wrapped[-1] == "":
        wrapped.pop()
    wrapped = wrapped or [""]
    chunks = [wrapped[i:i + rows] for i in range(0, len(wrapped), rows)]
    if len(chunks) > MAX_PAGES:
        notes.append(f"only the first {MAX_PAGES} pages were rendered")
        chunks = chunks[:MAX_PAGES]
    for n, ch in enumerate(chunks, 1):
        img = Image.new("RGB", (PAGE_W, PAGE_H), "white")
        d = ImageDraw.Draw(img)
        y = MARGIN
        for ln in ch:
            try:
                d.text((MARGIN, y), ln, fill=(25, 25, 25), font=font)
            except Exception:
                d.text((MARGIN, y), ln.encode("ascii", "replace").decode(), fill=(25, 25, 25), font=font)
            y += lh
        if len(chunks) > 1:
            d.text((PAGE_W - MARGIN - 70, PAGE_H - MARGIN // 2 - 10), f"{n}/{len(chunks)}",
                   fill=(150, 150, 150), font=small)
        yield img


def _reportlab_pdf(lines, out: Path) -> bool:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return False
    try:
        fname = "Courier"
        fp = _find_font(MONO_FONTS)
        if fp and fp.lower().endswith(".ttf"):
            try:
                pdfmetrics.registerFont(TTFont("FWMono", fp))
                fname = "FWMono"
            except Exception:
                pass
        size, lead, margin = 9, 12.5, 48
        W, H = A4
        cols = max(20, int((W - 2 * margin) / pdfmetrics.stringWidth("M", fname, size)))
        c = canvas.Canvas(str(out), pagesize=A4)
        c.setFont(fname, size)
        y = H - margin
        for ln in lines:
            for chunk in _wrap(ln, cols):
                if y < margin:
                    c.showPage()
                    c.setFont(fname, size)
                    y = H - margin
                c.drawString(margin, y, chunk)
                y -= lead
        c.save()
        return True
    except Exception as e:
        _log(f"reportlab failed: {e}")
        try:
            out.unlink()
        except Exception:
            pass
        return False


def _render_to_target(lines, out_dir: Path, stem: str, target: str, notes):
    if target == "PDF":
        out = _unique_path(out_dir, stem, "pdf")
        if _reportlab_pdf(lines, out):
            return [out]
    if not PIL_OK:
        raise RuntimeError("Pillow not installed (pip install Pillow)")
    return _save_pages(_render_pages(lines, notes), out_dir, stem, target, notes)


# ================================================================== writers
def _sheet_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "_", str(name)).strip("'")[:31] or "Sheet"
    base, i = t, 2
    while t.lower() in used:
        suf = f"_{i}"
        t = base[:31 - len(suf)] + suf
        i += 1
    used.add(t.lower())
    return t


def _put(ws, r, c, v):
    cell = ws.cell(row=r, column=c)
    if isinstance(v, str):
        cell.value = _clean(v)[:32000]
        cell.data_type = "s"  # never let "=..." text become a formula
    else:
        cell.value = v


def _write_xlsx(doc: Doc, out: Path, notes):
    if not XLSX_OK:
        raise RuntimeError("openpyxl not installed (pip install openpyxl)")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    tables = doc.tables
    imgs = [b for b in doc.blocks if b[0] == "img"]
    lines = doc.text_lines
    if any(l.strip() for l in lines) or not (tables or imgs):
        ws = wb.create_sheet(_sheet_title("Text", used))
        if len(lines) > MAX_XLSX_ROWS:
            notes.append("text truncated to Excel's row limit")
        for r, l in enumerate(lines[:MAX_XLSX_ROWS], 1):
            if l:
                _put(ws, r, 1, l)
        ws.column_dimensions["A"].width = 110
    for _, name, rows in tables:
        ws = wb.create_sheet(_sheet_title(name, used))
        if len(rows) > MAX_XLSX_ROWS:
            notes.append(f"sheet '{name}' truncated to Excel's row limit")
        widths = {}
        for r, row in enumerate(rows[:MAX_XLSX_ROWS], 1):
            for c, v in enumerate(row[:16384], 1):
                v = _auto_num(v)
                if v != "" and v is not None:
                    _put(ws, r, c, v)
                if r <= 200:
                    widths[c] = max(widths.get(c, 0), len(str(v)))
        for c, w in widths.items():
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = min(60, max(8, w + 2))
    for n, (_, path) in enumerate(imgs, 1):
        ws = wb.create_sheet(_sheet_title(f"Image {n}", used))
        ws.add_image(XLImage(path), "A1")
    wb.save(str(out))


def _write_docx(doc: Doc, out: Path):
    if not DOCX_OK:
        raise RuntimeError("python-docx not installed (pip install python-docx)")
    d = docx.Document()
    if doc.mono:
        st = d.styles["Normal"]
        st.font.name = "Consolas"
        st.font.size = Pt(9)
        st.paragraph_format.space_after = Pt(0)
    for b in doc.blocks:
        k = b[0]
        if k == "h":
            d.add_heading(_clean(b[1]), level=min(max(b[2], 0), 9))
        elif k == "p":
            d.add_paragraph(_clean(b[1]))
        elif k == "t":
            rows = b[2]
            if not rows:
                continue
            ncols = max(len(r) for r in rows) or 1
            t = d.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    t.cell(i, j).text = _clean("" if v is None else v)
            d.add_paragraph()
        elif k == "img":
            d.add_picture(b[1], width=Inches(6.0))
        elif k == "pb":
            d.add_page_break()
    d.save(str(out))


def _write_csv(doc: Doc, out_dir: Path, stem: str, notes):
    def dump(rows, path):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads UTF-8
            csv.writer(f).writerows([["" if c is None else c for c in r] for r in rows])

    tables = doc.tables
    if tables:
        if doc.has_text:
            notes.append("non-table text skipped (CSV keeps tables only)")
        if len(tables) == 1:
            out = _unique_path(out_dir, stem, "csv")
            dump(tables[0][2], out)
            return [out]
        outs = []
        for _, name, rows in tables:
            out = _unique_path(out_dir, f"{stem}_{name}", "csv")
            dump(rows, out)
            outs.append(out)
        notes.append(f"{len(outs)} tables saved as separate CSV files")
        return outs
    out = _unique_path(out_dir, stem, "csv")
    dump([[l] for l in doc.text_lines], out)
    return [out]


def _zip_paths(paths, out: Path, notes):
    used = set()

    def arc(name):
        base, i = name, 1
        while name in used:
            p = Path(base)
            name = str(p.with_name(f"{p.stem}_{i}{p.suffix}"))
            i += 1
        used.add(name)
        return name

    skip = out.resolve()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in map(Path, paths):
            if p.is_dir():
                for root, _, fs in os.walk(p):
                    for fn in fs:
                        fp = Path(root) / fn
                        if fp.resolve() == skip:
                            continue
                        try:
                            zf.write(fp, arc(str(Path(p.name) / fp.relative_to(p))))
                        except OSError as e:
                            notes.append(f"skipped {fn}: {e.strerror or e}")
            else:
                zf.write(p, arc(p.name))


_FONT_MAP = {"arial": "Arial", "helvetica": "Arial", "calibri": "Calibri", "cambria": "Cambria",
             "timesnewroman": "Times New Roman", "times": "Times New Roman", "courier": "Courier New",
             "couriernew": "Courier New", "verdana": "Verdana", "tahoma": "Tahoma",
             "georgia": "Georgia", "segoeui": "Segoe UI", "consolas": "Consolas"}


def _pdf_block_runs(b, W):
    lines, br = b["lines"], b["bbox"]
    full = all(l["bbox"][2] >= br[2] - 0.06 * W for l in lines[:-1])  # justified/full lines -> reflow
    runs, maxsize = [], 0.0
    for li, l in enumerate(lines):
        if li > 0:
            if not full:
                runs.append(("\n", False, False, 0, 10, ""))
            elif runs and not runs[-1][0].endswith(" "):
                runs.append((" ",) + runs[-1][1:])
        for s in l["spans"]:
            t = s["text"]
            if not t:
                continue
            name = s["font"].lower()
            bold = bool(s["flags"] & 16) or "bold" in name or "black" in name
            italic = bool(s["flags"] & 2) or "italic" in name or "oblique" in name
            runs.append((t, bold, italic, s["color"], s["size"], s["font"]))
            if t.strip():
                maxsize = max(maxsize, s["size"])
    return runs, maxsize


def _pdf_to_docx_rich(src: Path, out: Path, notes):
    """Rebuild a PDF as an editable DOCX: styled text, headings, tables, pictures and
    vector figures, in reading order."""
    d = fitz.open(str(src))
    try:
        if d.needs_pass and not d.authenticate(""):
            raise RuntimeError("PDF is password-protected")
        if len(d) == 0:
            raise RuntimeError("PDF has no pages")
        dicts = [pg.get_text("dict") for pg in d]
        sizes = {}
        for pd in dicts:
            for b in pd["blocks"]:
                if b.get("type") == 0:
                    for l in b["lines"]:
                        for s in l["spans"]:
                            n = len(s["text"].strip())
                            if n:
                                k = round(s["size"])
                                sizes[k] = sizes.get(k, 0) + n
        body = max(sizes, key=sizes.get) if sizes else 11
        doc = docx.Document()
        W0, H0 = d[0].rect.width, d[0].rect.height
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Pt(W0), Pt(H0)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(36)
        content_w = max(100.0, W0 - 72)
        zoom = fitz.Matrix(2.5, 2.5)
        scanned = False

        for pi, page in enumerate(d):
            pd, W, H = dicts[pi], page.rect.width, page.rect.height
            items, skip = [], []
            tblocks = [fitz.Rect(b["bbox"]) for b in pd["blocks"] if b["type"] == 0]

            try:  # tables
                for tb in page.find_tables().tables:
                    rows = [[("" if c is None else str(c)) for c in r] for r in tb.extract()]
                    if rows and max(len(r) for r in rows) > 1:
                        r = fitz.Rect(tb.bbox)
                        items.append((r.y0, r.x0, "t", rows))
                        skip.append(r)
            except Exception:
                pass

            for b in pd["blocks"]:  # raster pictures
                if b["type"] != 1:
                    continue
                r = fitz.Rect(b["bbox"])
                if r.width < 8 or r.height < 8 or r.width * r.height > 0.8 * W * H:
                    continue  # tiny bits / full-page backgrounds
                try:
                    png = page.get_pixmap(matrix=zoom, clip=r, alpha=False).tobytes("png")
                    items.append((r.y0, r.x0, "i", (png, r.width)))
                    skip.append(r)
                except Exception:
                    pass

            try:  # vector figures / charts
                for r in page.cluster_drawings():
                    r = fitz.Rect(r)
                    if r.width < 40 or r.height < 40 or r.width * r.height > 0.85 * W * H:
                        continue
                    if any((r & s).get_area() > 0.5 * r.get_area() for s in skip if r.intersects(s)):
                        continue
                    inside = sum(1 for t in tblocks if r.contains(fitz.Point((t.x0 + t.x1) / 2, (t.y0 + t.y1) / 2)))
                    if inside > 6:
                        continue  # decorative frame around real text
                    png = page.get_pixmap(matrix=zoom, clip=r, alpha=False).tobytes("png")
                    items.append((r.y0, r.x0, "i", (png, r.width)))
                    skip.append(r)
            except Exception:
                pass

            text_items = []
            for b in pd["blocks"]:
                if b["type"] != 0:
                    continue
                br = fitz.Rect(b["bbox"])
                ctr = fitz.Point((br.x0 + br.x1) / 2, (br.y0 + br.y1) / 2)
                if any(s.contains(ctr) for s in skip):
                    continue
                runs, maxsize = _pdf_block_runs(b, W)
                if not any(r[0].strip() for r in runs):
                    continue
                text_items.append((br, runs, maxsize))
            base = min((t[0].x0 for t in text_items), default=36)
            for br, runs, maxsize in text_items:
                items.append((br.y0, br.x0, "p", (br, runs, maxsize, base)))

            if not items:  # scanned page: OCR text if possible, else the page picture
                scanned = True
                pix = page.get_pixmap(matrix=zoom, alpha=False)
                if PIL_OK and ocr_available():
                    txt = _ocr(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
                    for para in re.split(r"\n\s*\n", txt.strip()):
                        if para.strip():
                            doc.add_paragraph(_clean(para.strip()))
                else:
                    doc.add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(content_w))

            for _, _, kind, data in sorted(items, key=lambda t: (round(t[0] / 4), t[1])):
                if kind == "t":
                    ncols = max(len(r) for r in data)
                    t = doc.add_table(rows=len(data), cols=ncols)
                    t.style = "Table Grid"
                    for i, row in enumerate(data):
                        for j, v in enumerate(row):
                            t.cell(i, j).text = _clean(v)
                    doc.add_paragraph()
                elif kind == "i":
                    doc.add_picture(io.BytesIO(data[0]), width=Pt(min(max(data[1], 20), content_w)))
                else:
                    br, runs, maxsize, base = data
                    plain = "".join(r[0] for r in runs)
                    ratio = maxsize / body if body else 1
                    all_bold = all(r[1] for r in runs if r[0].strip())
                    lvl = None
                    if len(plain) < 150 and "\n" not in plain:
                        lvl = 1 if ratio >= 1.7 else 2 if ratio >= 1.35 else 3 if (ratio >= 1.15 and all_bold) else None
                    para = doc.add_paragraph(style=f"Heading {lvl}") if lvl else doc.add_paragraph()
                    for t, bold, italic, color, sz, fname in runs:
                        if t == "\n":
                            para.add_run().add_break()
                            continue
                        r = para.add_run(_clean(t))
                        r.bold = True if bold else None
                        r.italic = True if italic else None
                        if not lvl:
                            r.font.size = Pt(max(6, round(sz * 2) / 2))
                        r.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
                        fam = re.split(r"[-,]", fname.split("+")[-1])[0].lower().replace(" ", "")
                        if fam in _FONT_MAP:
                            r.font.name = _FONT_MAP[fam]
                    centered = abs(br.x0 - (W - br.x1)) < 0.04 * W and br.width < 0.75 * W
                    if centered:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif br.x0 - base > 12:
                        para.paragraph_format.left_indent = Pt(min(br.x0 - base, content_w / 2))
                    para.paragraph_format.space_after = Pt(4)
            if pi < len(d) - 1:
                doc.add_page_break()
        if scanned:
            notes.append("scanned pages were OCR'd or inserted as images")
        doc.save(str(out))
    finally:
        d.close()


def _zip_many(files, notes):
    out = _unique_path(Path(CONFIG["output_dir"]), f"FormatWheel_{datetime.now():%Y%m%d_%H%M%S}", "zip")
    _zip_paths(files, out, notes)
    return str(out)


# ================================================================== the converter
def convert_file(src_path, target_fmt):
    """Convert one file. Returns (list_of_output_paths, notes). Raises RuntimeError if no route exists."""
    src = Path(src_path)
    target = str(target_fmt).upper()
    if target == "JPEG":
        target = "JPG"
    if target not in FORMATS:
        raise RuntimeError(f"unknown target format: {target_fmt}")
    if not src.exists():
        raise RuntimeError("source file not found")
    out_dir = Path(CONFIG["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    notes = []
    with tempfile.TemporaryDirectory(prefix="fwheel_") as tmp:
        outs = _convert(src, target, out_dir, Path(tmp), notes)
    return [str(o) for o in outs], list(dict.fromkeys(notes))


def _convert(src: Path, target: str, out_dir: Path, tmp: Path, notes):
    kind = _source_kind(str(src))
    ext = src.suffix.lower()
    tgt_ext = FORMATS[target]["ext"]

    if target == "ZIP":
        out = _unique_path(out_dir, src.stem if src.is_file() else src.name, "zip")
        _zip_paths([src], out, notes)
        return [out]
    if kind == "folder":
        raise RuntimeError("folders can only be converted to ZIP")
    if kind == "unknown":
        raise RuntimeError(f"unsupported file type ({ext or 'no extension'}) - only ZIP is possible")

    # --- Word/Excel -> PDF/images through a real Office engine when available
    if kind in ("docx", "xlsx") and target in PAGE_TARGETS:
        pdf = _office_convert(src, "pdf", tmp)
        if pdf:
            if target == "PDF":
                out = _unique_path(out_dir, src.stem, "pdf")
                shutil.copy2(pdf, out)
                return [out]
            pages = _pdf_page_images(pdf, notes)
            if pages is not None:
                return _save_pages(pages, out_dir, src.stem, target, notes)
        else:
            notes.append("MS Office/LibreOffice not found - used simplified rendering")

    # --- legacy formats: upgrade first
    if ext == ".doc":
        modern = _office_convert(src, "docx", tmp)
        if not modern:
            raise RuntimeError("legacy .doc needs MS Word or LibreOffice installed")
        src, ext = modern, ".docx"
    elif ext == ".xls":
        modern = _office_convert(src, "xlsx", tmp)
        if modern:
            src, ext = modern, ".xlsx"

    # --- already in that format: copy
    if _norm_ext(ext) == tgt_ext:
        out = _unique_path(out_dir, src.stem, tgt_ext)
        shutil.copy2(src, out)
        return [out]

    # --- images -> image/PDF
    if kind == "image" and target in PAGE_TARGETS:
        return _convert_image(src, target, out_dir, notes)

    # --- PDF -> images (real page renders)
    if kind == "pdf" and target in IMG_TARGETS:
        pages = _pdf_page_images(src, notes)
        if pages is not None:
            return _save_pages(pages, out_dir, src.stem, target, notes)
        notes.append("PyMuPDF not installed - rendered extracted text instead (pip install pymupdf)")

    # --- PDF -> DOCX: pdf2docx first, then our own PyMuPDF rebuild (styles, tables, pictures)
    if kind == "pdf" and target == "DOCX":
        out = _unique_path(out_dir, src.stem, "docx")
        done = False
        try:
            from pdf2docx import Converter
            cv = Converter(str(src))
            try:
                cv.convert(str(out))
            finally:
                cv.close()
            done = out.exists() and out.stat().st_size > 0
        except ImportError:
            notes.append("tip: pip install pdf2docx for even closer layout")
        except Exception as e:
            _log(f"pdf2docx failed: {e}")
        if not done and fitz is not None and DOCX_OK:
            try:
                _pdf_to_docx_rich(src, out, notes)
                done = True
            except RuntimeError:
                raise
            except Exception as e:
                _log(f"rich pdf->docx failed: {traceback.format_exc()}")
                notes.append("rich conversion failed - text-only DOCX created")
        if done:
            return [out]
        try:
            out.unlink()
        except Exception:
            pass

    # --- plain text -> TXT keeps the exact text
    if kind == "text" and target == "TXT":
        out = _unique_path(out_dir, src.stem, "txt")
        out.write_text(_read_text_any(src), encoding="utf-8")
        return [out]

    # --- generic route: source -> Doc -> target
    if kind == "text":
        doc = _read_text_doc(src, ext)
    elif kind == "docx":
        doc = _read_docx_doc(src, notes)
    elif kind == "xlsx":
        doc = _read_xlsx_doc(src, notes)
    elif kind == "pdf":
        doc = _read_pdf_doc(src, notes, target in ("XLSX", "CSV"), target == "DOCX", tmp)
    elif kind == "image":
        doc = _image_doc(src, target, tmp)
    else:
        raise RuntimeError(f"no conversion route for {kind} -> {target}")

    stem = src.stem
    if target == "TXT":
        out = _unique_path(out_dir, stem, "txt")
        out.write_text("\n".join(doc.lines()), encoding="utf-8")
        return [out]
    if target == "CSV":
        return _write_csv(doc, out_dir, stem, notes)
    if target == "XLSX":
        out = _unique_path(out_dir, stem, "xlsx")
        _write_xlsx(doc, out, notes)
        return [out]
    if target == "DOCX":
        out = _unique_path(out_dir, stem, "docx")
        _write_docx(doc, out)
        return [out]
    if target in PAGE_TARGETS:
        return _render_to_target(doc.lines(aligned=True), out_dir, stem, target, notes)
    raise RuntimeError(f"{kind} -> {target} not supported")


def _format_is_viable(path, target) -> bool:
    """Used to grey out wheel slices that cannot work for this file."""
    kind = _source_kind(str(path))
    ext = Path(path).suffix.lower()
    if target == "ZIP":
        return True
    if kind in ("folder", "unknown"):
        return False
    if _norm_ext(ext) == FORMATS[target]["ext"]:
        return False
    office = bool(WIN32_OK or _find_soffice())
    office_direct = kind in ("docx", "xlsx") and target in PAGE_TARGETS and office
    if kind == "image" and not PIL_OK:
        return False
    if kind == "docx":
        if ext == ".doc" and not office:
            return False
        if not DOCX_OK and not office_direct:
            return False
    if kind == "xlsx" and not XLSX_OK and not office_direct:
        return False
    if kind == "pdf" and not (fitz or pypdf):
        return False
    if target == "DOCX" and not DOCX_OK:
        return False
    if target == "XLSX" and not XLSX_OK:
        return False
    if target in IMG_TARGETS and not PIL_OK:
        return False
    if target == "PDF" and not PIL_OK and not _has_reportlab():
        return False
    if kind == "image" and target in ("TXT", "CSV") and not ocr_available():
        return False
    return True


# ================================================================== UI helpers
def _font(pt, bold=False, spacing=0.0):
    f = QFont()
    f.setFamilies(["Bahnschrift", "Segoe UI", "Arial"])
    f.setPointSizeF(pt)
    f.setBold(bold)
    if spacing:
        f.setLetterSpacing(QFont.SpacingType.AbsoluteSpacing, spacing)
    return f


def _alpha(c: QColor, a: int) -> QColor:
    d = QColor(c)
    d.setAlpha(max(0, min(255, int(a))))
    return d


def _force_foreground(hwnd: int):
    if sys.platform != "win32":
        return
    try:
        import ctypes
        u32, k32 = ctypes.windll.user32, ctypes.windll.kernel32
        fg = u32.GetForegroundWindow()
        fg_thread = u32.GetWindowThreadProcessId(fg, None)
        cur = k32.GetCurrentThreadId()
        u32.AttachThreadInput(fg_thread, cur, True)
        u32.SetForegroundWindow(hwnd)
        u32.BringWindowToTop(hwnd)
        u32.AttachThreadInput(fg_thread, cur, False)
    except Exception:
        pass


def _open_folder(path):
    try:
        if hasattr(os, "startfile"):
            os.startfile(str(path))
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception as e:
        _log(f"open folder failed: {e}")


# ================================================================== the wheel
class Slice:
    def __init__(self, key, angle, color, enabled):
        self.key, self.angle, self.color = key, angle, QColor(color)
        self.enabled, self.hover = enabled, 0.0


class RadialWheel(QWidget):
    chosen = pyqtSignal(str)
    closed = pyqtSignal()

    GAP = 1.6  # degrees between slices

    def __init__(self, files, theme_color="#00e5ff"):
        super().__init__()
        self.files = files
        self.theme = QColor(theme_color)
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint |
                            Qt.WindowType.WindowStaysOnTopHint | Qt.WindowType.Tool)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        self.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
        self.setMouseTracking(True)
        self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        self.setCursor(Qt.CursorShape.BlankCursor)

        screen = QApplication.screenAt(QCursor.pos()) or QApplication.primaryScreen()
        geo = screen.geometry()
        self.setGeometry(geo)

        self.r_out = int(max(160, min(420, CONFIG["wheel_radius"])))
        self.r_in = int(self.r_out * 0.46)
        margin = self.r_out + 80
        cur = QCursor.pos() - geo.topLeft()
        cx = min(max(cur.x(), margin), geo.width() - margin) if geo.width() > 2 * margin else geo.width() / 2
        cy = min(max(cur.y(), margin), geo.height() - margin) if geo.height() > 2 * margin else geo.height() / 2
        self.center = QPointF(cx, cy)
        self._wheel_rect = QRect(int(cx - margin), int(cy - margin), 2 * margin, 2 * margin)
        QCursor.setPos(geo.topLeft() + QPoint(int(cx), int(cy)))  # aim from the middle, like a game

        n = len(WHEEL_ORDER)
        self.step = 360.0 / n
        self.slices = []
        for i, k in enumerate(WHEEL_ORDER):
            ok = any(_format_is_viable(f, k) for f in files)
            self.slices.append(Slice(k, -90 + i * self.step, FORMATS[k]["color"], ok))

        names = [Path(f).name for f in files]
        self.summary = names[0] + (f"   +{len(names) - 1} more" if len(names) > 1 else "")
        self.sel = None
        self.mouse = QPointF(self.center)
        self._kb = False
        self._closing = False
        self._pending = None
        self.opacity, self.scale = 0.0, 0.82

        self._timer = QTimer(self)
        self._timer.timeout.connect(self._tick)
        self._timer.start(16)

    # ---- lifecycle
    def showEvent(self, e):
        super().showEvent(e)
        QTimer.singleShot(0, self._grab_focus)

    def _grab_focus(self):
        self.raise_()
        self.activateWindow()
        self.setFocus()
        _force_foreground(int(self.winId()))

    def closeEvent(self, e):
        self._timer.stop()
        self.closed.emit()
        super().closeEvent(e)

    def confirm(self):
        if self._closing or self.sel is None or not self.slices[self.sel].enabled:
            return
        key = self.slices[self.sel].key
        self._closing = True
        self._pending = lambda: self.chosen.emit(key)

    def cancel(self):
        self._closing = True
        self._pending = None

    def on_hotkey_release(self):
        if CONFIG.get("release_to_select", True):
            self.confirm()

    def _finish(self):
        cb = self._pending
        self.close()
        if cb:
            cb()

    # ---- animation
    def _repaint(self):
        if self.opacity < 1 or self.scale < 1 or self._closing:
            self.update()
        else:
            self.update(self._wheel_rect)

    def _tick(self):
        dirty = False
        if self._closing:
            self.opacity = max(0.0, self.opacity - 0.16)
            self.scale += 0.02
            if self.opacity <= 0:
                self._finish()
                return
            dirty = True
        else:
            if self.opacity < 1:
                self.opacity = min(1.0, self.opacity + 0.14)
                dirty = True
            if self.scale < 1:
                self.scale = min(1.0, self.scale + (1 - self.scale) * 0.28 + 0.004)
                dirty = True
        for i, s in enumerate(self.slices):
            tgt = 1.0 if i == self.sel else 0.0
            d = tgt - s.hover
            if abs(d) > 0.01:
                s.hover += d * 0.32
                dirty = True
            elif s.hover != tgt:
                s.hover = tgt
                dirty = True
        if dirty:
            self._repaint()

    # ---- selection
    def _index_at(self, pos: QPointF):
        dx, dy = pos.x() - self.center.x(), pos.y() - self.center.y()
        dist = math.hypot(dx, dy)
        if dist < self.r_in * 0.35 or dist > self.r_out + 160:
            return None
        ang = math.degrees(math.atan2(dy, dx))
        return int(((ang + 90 + self.step / 2) % 360) // self.step) % len(self.slices)

    def _cycle(self, step):
        n = len(self.slices)
        i = self.sel if self.sel is not None else (-1 if step > 0 else 0)
        for _ in range(n):
            i = (i + step) % n
            if self.slices[i].enabled:
                self._kb, self.sel = True, i
                self._repaint()
                return

    def mouseMoveEvent(self, e):
        self._kb = False
        self.mouse = e.position()
        self.sel = self._index_at(self.mouse)
        self._repaint()

    def mousePressEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            if self.sel is not None and self.slices[self.sel].enabled:
                self.confirm()
            elif self.sel is None:
                self.cancel()
        else:
            self.cancel()

    def wheelEvent(self, e):
        self._cycle(1 if e.angleDelta().y() < 0 else -1)

    def keyPressEvent(self, e):
        k = e.key()
        if k in (Qt.Key.Key_Escape, Qt.Key.Key_Backspace):
            self.cancel()
        elif k in (Qt.Key.Key_Return, Qt.Key.Key_Enter, Qt.Key.Key_Space):
            self.confirm()
        elif k in (Qt.Key.Key_Left, Qt.Key.Key_Up):
            self._cycle(-1)
        elif k in (Qt.Key.Key_Right, Qt.Key.Key_Down, Qt.Key.Key_Tab):
            self._cycle(1)

    # ---- painting
    def _slice_path(self, angle, ri, ro):
        c = self.center
        h = self.step / 2 - self.GAP
        outer = QRectF(c.x() - ro, c.y() - ro, 2 * ro, 2 * ro)
        inner = QRectF(c.x() - ri, c.y() - ri, 2 * ri, 2 * ri)
        start, span = -angle - h, 2 * h
        path = QPainterPath()
        path.arcMoveTo(outer, start)
        path.arcTo(outer, start, span)
        path.arcTo(inner, start + span, -span)
        path.closeSubpath()
        return path

    def _pointer_pos(self):
        c = self.center
        if self._kb and self.sel is not None:
            a = math.radians(self.slices[self.sel].angle)
            r = self.r_in * 0.62
            return QPointF(c.x() + r * math.cos(a), c.y() + r * math.sin(a))
        vx, vy = self.mouse.x() - c.x(), self.mouse.y() - c.y()
        L, lim = math.hypot(vx, vy), self.r_out - 10
        if L > lim and L > 0:
            vx, vy = vx * lim / L, vy * lim / L
        return QPointF(c.x() + vx, c.y() + vy)

    def _draw_icon(self, p, x, y, sc, color, label, enabled, hv):
        w, h, fold = 30 * sc, 38 * sc, 9 * sc
        l, t = x - w / 2, y - h / 2 - 9 * sc
        page = QPainterPath()
        page.moveTo(l, t)
        page.lineTo(l + w - fold, t)
        page.lineTo(l + w, t + fold)
        page.lineTo(l + w, t + h)
        page.lineTo(l, t + h)
        page.closeSubpath()
        col = QColor(color) if enabled else QColor(88, 94, 106)
        g = QLinearGradient(l, t, l, t + h)
        g.setColorAt(0, col.lighter(125))
        g.setColorAt(1, col.darker(150))
        p.setBrush(QBrush(g))
        p.setPen(QPen(QColor(255, 255, 255, int(110 + 120 * hv) if enabled else 50), 1.3))
        p.drawPath(page)
        tri = QPainterPath()
        tri.moveTo(l + w - fold, t)
        tri.lineTo(l + w - fold, t + fold)
        tri.lineTo(l + w, t + fold)
        tri.closeSubpath()
        p.setBrush(QColor(255, 255, 255, 120 if enabled else 40))
        p.setPen(Qt.PenStyle.NoPen)
        p.drawPath(tri)
        p.setPen(QPen(QColor(255, 255, 255, 110 if enabled else 35), 1.4 * sc))
        for k in range(3):
            yy = t + 17 * sc + k * 6 * sc
            p.drawLine(QPointF(l + 6 * sc, yy), QPointF(l + w - 6 * sc, yy))
        p.setPen(QColor(255, 255, 255) if enabled else QColor(115, 122, 136))
        p.setFont(_font(11 + 2.5 * hv, True, 1.2))
        p.drawText(QRectF(x - 50, t + h + 3 * sc, 100, 22), Qt.AlignmentFlag.AlignCenter, label)

    def paintEvent(self, _):
        p = QPainter(self)
        p.setRenderHints(QPainter.RenderHint.Antialiasing | QPainter.RenderHint.TextAntialiasing)
        W, H = self.width(), self.height()
        cx, cy = self.center.x(), self.center.y()
        op = self.opacity
        dim = max(40, min(230, int(CONFIG.get("backdrop_dim", 150))))

        # backdrop (vignette + theme haze)
        vg = QRadialGradient(self.center, max(W, H) * 0.75)
        vg.setColorAt(0.0, QColor(0, 0, 0, int(dim * 0.55 * op) + 1))
        vg.setColorAt(1.0, QColor(0, 0, 0, int(dim * op) + 1))
        p.fillRect(self.rect(), QBrush(vg))
        haze = QRadialGradient(self.center, self.r_out * 1.7)
        haze.setColorAt(0.5, _alpha(self.theme, 38 * op))
        haze.setColorAt(1.0, QColor(0, 0, 0, 0))
        p.setPen(Qt.PenStyle.NoPen)
        p.setBrush(QBrush(haze))
        p.drawEllipse(self.center, self.r_out * 1.7, self.r_out * 1.7)

        p.setOpacity(op)
        p.translate(self.center)
        p.scale(self.scale, self.scale)
        p.translate(-self.center.x(), -self.center.y())

        # HUD tick ring
        rt = self.r_out + 26
        for i in range(120):
            a = math.radians(i * 3)
            ln = 9 if i % 10 == 0 else 4
            p.setPen(QPen(_alpha(self.theme, 140 if i % 10 == 0 else 55), 1.2))
            p.drawLine(QPointF(cx + rt * math.cos(a), cy + rt * math.sin(a)),
                       QPointF(cx + (rt + ln) * math.cos(a), cy + (rt + ln) * math.sin(a)))

        # slices (hovered one drawn last so it overlaps)
        order = sorted(range(len(self.slices)), key=lambda i: self.slices[i].hover)
        for i in order:
            s = self.slices[i]
            hv = s.hover
            ro, ri = self.r_out + 16 * hv, self.r_in
            path = self._slice_path(s.angle, ri, ro)
            if s.enabled and hv > 0.02:
                p.setBrush(Qt.BrushStyle.NoBrush)
                p.setPen(QPen(_alpha(s.color, 70 * hv), 12))
                p.drawPath(path)
            p.setBrush(QColor(12, 14, 20, 190 if s.enabled else 150))
            p.setPen(QPen(QColor(255, 255, 255, 42 if s.enabled else 16), 1.2))
            p.drawPath(path)
            if hv > 0.02:
                g = QRadialGradient(self.center, ro)
                g.setColorAt(max(0.0, min(0.95, ri / ro)), _alpha(s.color, (45 if s.enabled else 10) * hv))
                g.setColorAt(1.0, _alpha(s.color, (215 if s.enabled else 30) * hv))
                p.setBrush(QBrush(g))
                p.setPen(QPen(QColor(255, 255, 255, int(255 * hv)), 2.6))
                p.drawPath(path)
            a = math.radians(s.angle)
            rm = (ri + ro) / 2
            self._draw_icon(p, cx + rm * math.cos(a), cy + rm * math.sin(a),
                            1.0 + 0.24 * hv, s.color, s.key, s.enabled, hv)

        # centre disc
        sel = self.slices[self.sel] if self.sel is not None else None
        disc = self.r_in - 12
        ring = _alpha(sel.color, 230) if sel and sel.enabled else _alpha(self.theme, 150)
        p.setBrush(QColor(8, 10, 15, 228))
        p.setPen(QPen(ring, 2.2))
        p.drawEllipse(self.center, disc, disc)
        p.setPen(QPen(_alpha(ring, 70), 1))
        p.setBrush(Qt.BrushStyle.NoBrush)
        p.drawEllipse(self.center, disc - 9, disc - 9)
        box = lambda dy, h: QRectF(cx - disc + 10, cy + dy, 2 * disc - 20, h)
        center_flag = Qt.AlignmentFlag.AlignCenter
        if sel is None:
            p.setPen(QColor(235, 242, 255))
            p.setFont(_font(21, True, 3))
            p.drawText(box(-40, 32), center_flag, "FORMAT")
            p.setPen(self.theme)
            p.drawText(box(-10, 32), center_flag, "WHEEL")
            p.setPen(QColor(130, 146, 172))
            p.setFont(_font(8.5))
            p.drawText(box(28, 18), center_flag, "move to choose")
        else:
            p.setPen(QColor(255, 255, 255) if sel.enabled else QColor(120, 128, 142))
            p.setFont(_font(34, True, 2))
            p.drawText(box(-42, 46), center_flag, sel.key)
            p.setFont(_font(9.5, True, 1))
            if sel.enabled:
                p.setPen(_alpha(sel.color.lighter(130), 255))
                p.drawText(box(8, 20), center_flag, FORMATS[sel.key]["desc"].upper())
                p.setPen(QColor(150, 164, 190))
                p.setFont(_font(8.5))
                n = len(self.files)
                p.drawText(box(30, 18), center_flag, f"{n} file{'s' if n != 1 else ''}")
            else:
                p.setPen(QColor(255, 99, 99))
                p.drawText(box(8, 20), center_flag, "UNAVAILABLE")
                p.setPen(QColor(150, 164, 190))
                p.setFont(_font(8.5))
                p.drawText(box(30, 18), center_flag, "not possible here")

        # pointer (our own cursor, like a game stick indicator)
        pt = self._pointer_pos()
        p.setPen(QPen(QColor(255, 255, 255, 60), 1))
        p.drawLine(self.center, pt)
        p.setBrush(QColor(255, 255, 255, 240))
        p.setPen(QPen(QColor(0, 0, 0, 160), 1.5))
        p.drawEllipse(pt, 5.5, 5.5)

        # captions
        p.setFont(_font(10.5, True, 2))
        p.setPen(_alpha(self.theme, 230))
        fm = QFontMetrics(p.font())
        p.drawText(QRectF(cx - 330, cy - self.r_out - 62, 660, 24), center_flag,
                   fm.elidedText(self.summary.upper(), Qt.TextElideMode.ElideMiddle, 640))
        p.setFont(_font(8.5, False, 1))
        p.setPen(QColor(160, 174, 198, 200))
        p.drawText(QRectF(cx - 330, cy + self.r_out + 44, 660, 20), center_flag,
                   "CLICK / ENTER  select      ESC / RIGHT-CLICK  cancel      SCROLL  cycle")
        p.end()


# ================================================================== toast notifications
class Toast(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint | Qt.WindowType.WindowStaysOnTopHint |
                            Qt.WindowType.Tool | Qt.WindowType.WindowDoesNotAcceptFocus)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        self.setAttribute(Qt.WidgetAttribute.WA_ShowWithoutActivating)
        self.title, self.lines = "", []
        self.accent = QColor("#00e5ff")
        self.busy, self.on_click = False, None
        self.alpha, self.target, self.phase = 0.0, 0.0, 0.0
        self.timer = QTimer(self)
        self.timer.timeout.connect(self._tick)
        self.life = QTimer(self)
        self.life.setSingleShot(True)
        self.life.timeout.connect(self.dismiss)

    def popup(self, title, lines=(), accent="#00e5ff", duration=4500, busy=False, on_click=None):
        self.title, self.lines = title, list(lines)[:5]
        self.accent, self.busy, self.on_click = QColor(accent), busy, on_click
        w = 460
        h = 56 + 20 * len(self.lines) + (8 if self.lines else 0)
        self.setFixedSize(w, h)
        scr = QApplication.primaryScreen().availableGeometry()
        self.move(scr.center().x() - w // 2, scr.bottom() - h - 36)
        self.target = 1.0
        self.show()
        self.raise_()
        self.timer.start(16)
        self.life.stop()
        if duration > 0 and not busy:
            self.life.start(duration)

    def dismiss(self):
        self.target = 0.0
        self.timer.start(16)

    def _tick(self):
        self.phase = (self.phase + 0.02) % 1.0
        if self.alpha < self.target:
            self.alpha = min(self.target, self.alpha + 0.12)
        elif self.alpha > self.target:
            self.alpha = max(self.target, self.alpha - 0.10)
        self.setWindowOpacity(self.alpha)
        self.update()
        if self.alpha <= 0 and self.target == 0:
            self.timer.stop()
            self.hide()
        elif self.alpha >= 1 and not self.busy:
            self.timer.stop()

    def mousePressEvent(self, e):
        if self.on_click:
            try:
                self.on_click()
            except Exception:
                pass
        self.dismiss()

    def paintEvent(self, _):
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        r = QRectF(self.rect()).adjusted(1, 1, -1, -1)
        path = QPainterPath()
        path.addRoundedRect(r, 10, 10)
        p.fillPath(path, QBrush(QColor(14, 16, 22, 240)))
        p.setPen(QPen(QColor(255, 255, 255, 34), 1))
        p.drawPath(path)
        p.save()
        p.setClipPath(path)
        p.fillRect(QRectF(0, 0, 6, self.height()), self.accent)
        if self.busy:
            seg = 130
            x = self.phase * (self.width() + seg) - seg
            p.fillRect(QRectF(x, self.height() - 3, seg, 3), self.accent)
        p.restore()
        w = self.width()
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(13, True, 0.6))
        p.drawText(QRectF(22, 12, w - 36, 28), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                   self.title)
        p.setFont(_font(9.5))
        p.setPen(QColor(172, 184, 204))
        fm = QFontMetrics(p.font())
        for i, l in enumerate(self.lines):
            p.drawText(QRectF(22, 44 + i * 20, w - 36, 20),
                       Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                       fm.elidedText(l, Qt.TextElideMode.ElideMiddle, w - 40))
        p.end()


# ================================================================== settings dialog
class SettingsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Format Wheel - Settings")
        self.setMinimumWidth(480)
        self.setStyleSheet("""
            QDialog { background:#1a1d26; color:#e0e6f0; }
            QLabel { color:#c8d2e0; }
            QLineEdit { background:#252a36; color:#e0e6f0; border:1px solid #3a4150;
                border-radius:6px; padding:6px; }
            QPushButton { background:#2f6bff; color:white; border-radius:6px;
                padding:8px 16px; font-weight:600; }
            QPushButton:hover { background:#4a7fff; }
            QCheckBox { color:#c8d2e0; }
        """)
        lay = QVBoxLayout(self)
        lay.addWidget(QLabel("<h2 style='color:#00e5ff'>Format Wheel Settings</h2>"))

        lay.addWidget(QLabel("Activation hotkey (e.g. ctrl+alt+w):"))
        self.hotkey_edit = QLineEdit(CONFIG["hotkey"])
        lay.addWidget(self.hotkey_edit)

        lay.addWidget(QLabel("Output folder:"))
        row = QHBoxLayout()
        self.out_edit = QLineEdit(CONFIG["output_dir"])
        row.addWidget(self.out_edit)
        b = QPushButton("Browse...")
        b.clicked.connect(self._browse)
        row.addWidget(b)
        lay.addLayout(row)

        lay.addWidget(QLabel("Theme color:"))
        trow = QHBoxLayout()
        self.theme_edit = QLineEdit(CONFIG["theme_color"])
        trow.addWidget(self.theme_edit)
        cb = QPushButton("Pick...")
        cb.clicked.connect(self._pick_color)
        trow.addWidget(cb)
        lay.addLayout(trow)

        lay.addWidget(QLabel("Wheel radius (px, 160-420):"))
        self.rad_edit = QLineEdit(str(CONFIG["wheel_radius"]))
        lay.addWidget(self.rad_edit)

        lay.addWidget(QLabel("Screen dimming behind the wheel (40-230):"))
        self.dim_edit = QLineEdit(str(CONFIG["backdrop_dim"]))
        lay.addWidget(self.dim_edit)

        def chk(text, key):
            c = QCheckBox(text)
            c.setChecked(bool(CONFIG[key]))
            lay.addWidget(c)
            return c

        self.release_chk = chk("Hold hotkey, aim, release to select (GTA style)", "release_to_select")
        self.start_win = chk("Start with Windows", "start_with_windows")
        self.clip_chk = chk("Allow converting clipboard image when no file selected", "auto_convert_clipboard")
        self.copy_chk = chk("Copy result to clipboard after conversion", "copy_after_convert")
        self.open_chk = chk("Open output folder after conversion", "open_folder_after")
        self.quiet_chk = chk("Do not show success notifications", "quiet_success")

        lay.addSpacing(10)
        btn_row = QHBoxLayout()
        btn_row.addStretch()
        cancel = QPushButton("Cancel")
        cancel.clicked.connect(self.reject)
        save = QPushButton("Save")
        save.clicked.connect(self._save)
        btn_row.addWidget(cancel)
        btn_row.addWidget(save)
        lay.addLayout(btn_row)

    def _browse(self):
        d = QFileDialog.getExistingDirectory(self, "Choose output folder", self.out_edit.text())
        if d:
            self.out_edit.setText(d)

    def _pick_color(self):
        c = QColorDialog.getColor(QColor(self.theme_edit.text()), self, "Theme color")
        if c.isValid():
            self.theme_edit.setText(c.name())

    def _save(self):
        hk = self.hotkey_edit.text().strip().lower()
        if hk:
            CONFIG["hotkey"] = hk
        CONFIG["output_dir"] = self.out_edit.text().strip() or DEFAULT_CONFIG["output_dir"]
        col = self.theme_edit.text().strip()
        if QColor(col).isValid():
            CONFIG["theme_color"] = col
        for key, edit, lo, hi in (("wheel_radius", self.rad_edit, 160, 420),
                                  ("backdrop_dim", self.dim_edit, 40, 230)):
            try:
                CONFIG[key] = max(lo, min(hi, int(edit.text())))
            except ValueError:
                pass
        CONFIG["release_to_select"] = self.release_chk.isChecked()
        CONFIG["start_with_windows"] = self.start_win.isChecked()
        CONFIG["auto_convert_clipboard"] = self.clip_chk.isChecked()
        CONFIG["copy_after_convert"] = self.copy_chk.isChecked()
        CONFIG["open_folder_after"] = self.open_chk.isChecked()
        CONFIG["quiet_success"] = self.quiet_chk.isChecked()
        save_config(CONFIG)
        try:
            Path(CONFIG["output_dir"]).mkdir(parents=True, exist_ok=True)
        except Exception as e:
            _log(f"cannot create output dir: {e}")
        apply_startup()
        hotkey_manager.restart()
        self.accept()


def apply_startup():
    if sys.platform != "win32":
        return
    try:
        import winreg
        key_path = r"Software\Microsoft\Windows\CurrentVersion\Run"
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path, 0, winreg.KEY_SET_VALUE) as k:
            if CONFIG["start_with_windows"]:
                if getattr(sys, "frozen", False):
                    cmd = f'"{sys.executable}"'
                else:
                    exe = Path(sys.executable)
                    pyw = exe.with_name("pythonw.exe")
                    exe = pyw if pyw.exists() else exe
                    cmd = f'"{exe}" "{Path(__file__).resolve()}"'
                winreg.SetValueEx(k, "FormatWheel", 0, winreg.REG_SZ, cmd)
            else:
                try:
                    winreg.DeleteValue(k, "FormatWheel")
                except FileNotFoundError:
                    pass
    except Exception as e:
        _log(f"startup registry: {e}")


# ================================================================== global hotkey
class HotkeyBridge(QObject):
    fired = pyqtSignal()
    released = pyqtSignal()


bridge = HotkeyBridge()

_KEY_ALIASES = {"control": "ctrl", "ctl": "ctrl", "windows": "win", "super": "win", "cmd": "win",
                "return": "enter", "escape": "esc"}


class HotkeyManager:
    def __init__(self):
        self._listener = None
        self._pressed = set()
        self._fired = False

    @staticmethod
    def _parse_combo(combo):
        return {_KEY_ALIASES.get(k.strip().lower(), k.strip().lower())
                for k in combo.split("+") if k.strip()}

    def _start_listener(self):
        try:
            from pynput import keyboard
        except ImportError:
            _log("pynput not installed - hotkey disabled")
            print("pynput not installed - hotkey disabled. pip install pynput")
            return
        target = self._parse_combo(CONFIG["hotkey"])
        self._pressed, self._fired = set(), False
        names = {"ctrl_l": "ctrl", "ctrl_r": "ctrl", "shift_l": "shift", "shift_r": "shift",
                 "alt_l": "alt", "alt_r": "alt", "cmd": "win", "cmd_l": "win", "cmd_r": "win"}

        def norm(k):
            if isinstance(k, keyboard.Key):
                return names.get(k.name, k.name)
            vk = getattr(k, "vk", None)  # char is unreliable while Ctrl is held
            if vk is not None:
                if 65 <= vk <= 90:
                    return chr(vk).lower()
                if 48 <= vk <= 57:
                    return chr(vk)
                if 96 <= vk <= 105:
                    return str(vk - 96)
            ch = getattr(k, "char", None)
            return ch.lower() if ch and ch.isprintable() else None

        def on_press(k):
            n = norm(k)
            if not n:
                return
            self._pressed.add(n)
            if target and target.issubset(self._pressed) and not self._fired:
                self._fired = True
                bridge.fired.emit()

        def on_release(k):
            n = norm(k)
            if n:
                self._pressed.discard(n)
            if n in target and self._fired:
                self._fired = False
                bridge.released.emit()

        self._listener = keyboard.Listener(on_press=on_press, on_release=on_release)
        self._listener.daemon = True
        self._listener.start()

    def stop(self):
        if self._listener is not None:
            try:
                self._listener.stop()
            except Exception:
                pass
            self._listener = None

    def restart(self):
        self.stop()
        self._start_listener()


hotkey_manager = HotkeyManager()


# ================================================================== selection capture
def _files_from_explorer():
    if not WIN32_OK:
        return []
    try:
        import ctypes
        fg = int(ctypes.windll.user32.GetForegroundWindow())
        shell = win32com.client.Dispatch("Shell.Application")
        candidates = []
        for w in shell.Windows():
            try:
                if "explorer" not in w.FullName.lower():
                    continue
                files = [it.Path for it in w.Document.SelectedItems()
                         if os.path.isfile(it.Path) or os.path.isdir(it.Path)]
                if files:
                    candidates.append((int(w.HWND), files))
            except Exception:
                continue
        for hwnd, files in candidates:
            if hwnd == fg:
                return files
        return candidates[-1][1] if candidates else []
    except Exception as e:
        _log(f"explorer enum: {e}")
        return []


def _files_from_clipboard():
    files = []
    try:
        md = QApplication.clipboard().mimeData()
        if md and md.hasUrls():
            for u in md.urls():
                if u.isLocalFile():
                    p = u.toLocalFile()
                    if os.path.exists(p):
                        files.append(p)
    except Exception:
        pass
    return files


CLIP_DIR = Path(tempfile.gettempdir()) / "FormatWheel_clip"


def _clipboard_image_to_temp():
    if not CONFIG["auto_convert_clipboard"]:
        return []
    try:
        img = QApplication.clipboard().image()
        if not img.isNull():
            CLIP_DIR.mkdir(exist_ok=True)
            tmp = CLIP_DIR / f"Clipboard_{datetime.now():%Y%m%d_%H%M%S}.png"
            img.save(str(tmp), "PNG")
            return [str(tmp)]
    except Exception as e:
        _log(f"clipboard image: {e}")
    return []


def get_selected_files():
    return _files_from_explorer() or _files_from_clipboard() or _clipboard_image_to_temp()


# ================================================================== app glue
class App(QApplication):
    wheel = None
    toast = None
    worker = None

    def show_settings(self):
        SettingsDialog().exec()


def _toast(title, lines=(), kind="info", duration=4500, busy=False, on_click=None):
    accent = {"info": CONFIG["theme_color"], "ok": "#34c759", "warn": "#ff9500", "err": "#ff3b30"}[kind]
    if App.toast is not None:
        App.toast.popup(title, lines, accent, duration, busy, on_click)


class ConvertThread(QThread):
    progress = pyqtSignal(int, int, str)
    done = pyqtSignal(list, list, list, str, list)

    def __init__(self, files, fmt):
        super().__init__()
        self.files, self.fmt = list(files), fmt

    def run(self):
        co = False
        try:
            pythoncom.CoInitialize()
            co = True
        except Exception:
            pass
        results, errors, notes = [], [], []
        try:
            if self.fmt == "ZIP" and len(self.files) > 1:
                try:
                    results.append(_zip_many(self.files, notes))
                except Exception as e:
                    _log(traceback.format_exc())
                    errors.append(f"ZIP: {e}")
            else:
                multi = len(self.files) > 1
                for i, f in enumerate(self.files, 1):
                    name = Path(f).name
                    self.progress.emit(i, len(self.files), name)
                    try:
                        outs, nt = convert_file(f, self.fmt)
                        ok = [o for o in outs if Path(o).exists()]
                        if ok:
                            results.extend(ok)
                            notes.extend((f"{name}: {n}" if multi else n) for n in nt)
                        else:
                            errors.append(f"{name}: output file was not created")
                    except Exception as e:
                        _log(f"{name} -> {self.fmt}\n{traceback.format_exc()}")
                        errors.append(f"{name}: {e}")
        finally:
            if co:
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
        self.done.emit(results, errors, list(dict.fromkeys(notes)), self.fmt, self.files)


def _on_progress(i, n, name):
    _toast(f"Converting {i}/{n}", [name], busy=True)


def _on_done(results, errors, notes, fmt, files):
    try:
        if CLIP_DIR.exists() and any(str(CLIP_DIR) in str(f) for f in files):
            for f in CLIP_DIR.glob("Clipboard_*.png"):
                try:
                    f.unlink()
                except Exception:
                    pass
    except Exception:
        pass
    out_dir = Path(CONFIG["output_dir"])
    if CONFIG["copy_after_convert"] and results:
        try:
            md = QMimeData()
            md.setUrls([QUrl.fromLocalFile(r) for r in results])
            QApplication.clipboard().setMimeData(md)
        except Exception as e:
            _log(f"clipboard copy failed: {e}")
    if results and CONFIG["open_folder_after"]:
        _open_folder(out_dir)
    opener = lambda: _open_folder(out_dir)
    if errors and not results:
        _toast(f"Conversion to {fmt} failed", errors[:4], "err", 9000)
    elif errors:
        _toast(f"{len(results)} converted, {len(errors)} failed",
               errors[:3] + notes[:1], "warn", 9000, on_click=opener)
    elif CONFIG["quiet_success"] and not notes:
        if App.toast:
            App.toast.dismiss()
    else:
        _toast(f"Converted to {fmt} - {len(results)} file{'s' if len(results) != 1 else ''}",
               [f"Saved to {out_dir}"] + notes[:3] + ["Click to open folder"], "ok", 5500, on_click=opener)


def _do_convert(files, fmt):
    if App.worker is not None and App.worker.isRunning():
        _toast("Busy", ["Still converting the previous batch"], "warn", 3000)
        return
    th = ConvertThread(files, fmt)
    App.worker = th
    th.progress.connect(_on_progress)
    th.done.connect(_on_done)
    _toast(f"Converting to {fmt}...", [Path(files[0]).name], busy=True)
    th.start()


def show_wheel(files):
    w = RadialWheel(files, CONFIG["theme_color"])
    w.chosen.connect(lambda fmt: _do_convert(files, fmt))
    w.closed.connect(lambda: setattr(App, "wheel", None) if App.wheel is w else None)
    App.wheel = w
    w.show()


def show_wheel_from_selection():
    try:
        if App.wheel is not None:
            App.wheel.cancel()
            return
        files = get_selected_files()
        if not files:
            _toast("Nothing selected",
                   ["Select files in Explorer (or copy them / a screenshot) and try again"],
                   "warn", 3500)
            return
        show_wheel(files)
    except Exception as e:
        _log(traceback.format_exc())
        _toast("Format Wheel error", [str(e)], "err", 7000)


def _on_hotkey_release():
    if App.wheel is not None:
        App.wheel.on_hotkey_release()


def _excepthook(etype, value, tb):
    _log("UNHANDLED: " + "".join(traceback.format_exception(etype, value, tb)))
    try:
        _toast("Unexpected error", [str(value), f"Details: {LOG_FILE}"], "err", 8000)
    except Exception:
        pass


def make_icon(color="#00e5ff"):
    pm = QPixmap(64, 64)
    pm.fill(Qt.GlobalColor.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.RenderHint.Antialiasing)
    c = QColor(color)
    p.setBrush(QBrush(c))
    p.setPen(QPen(c.lighter(150), 2))
    p.drawEllipse(6, 6, 52, 52)
    p.setBrush(QColor(20, 22, 30))
    p.drawEllipse(20, 20, 24, 24)
    p.setPen(QColor(c))
    p.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
    p.drawText(QRectF(0, 0, 64, 64), Qt.AlignmentFlag.AlignCenter, "\u21c4")
    p.end()
    return QIcon(pm)


def main():
    sys.excepthook = _excepthook  # PyQt6 aborts on unhandled slot exceptions otherwise
    app = App(sys.argv)
    app.setQuitOnLastWindowClosed(False)
    app.setStyle("Fusion")

    lock = QLockFile(str(Path(tempfile.gettempdir()) / "FormatWheel.lock"))
    lock.setStaleLockTime(5000)
    if not lock.tryLock(100):
        print("Format Wheel is already running.")
        return

    App.toast = Toast()
    bridge.fired.connect(show_wheel_from_selection)
    bridge.released.connect(_on_hotkey_release)

    tray = QSystemTrayIcon(make_icon(CONFIG["theme_color"]), app)
    tray.setToolTip(f"Format Wheel ({CONFIG['hotkey']})")
    menu = QMenu()
    a1 = QAction("Convert selected files...", menu)
    a1.triggered.connect(show_wheel_from_selection)
    menu.addAction(a1)
    a2 = QAction("Open output folder", menu)
    a2.triggered.connect(lambda: _open_folder(CONFIG["output_dir"]))
    menu.addAction(a2)
    a3 = QAction("Settings...", menu)
    a3.triggered.connect(app.show_settings)
    menu.addAction(a3)
    menu.addSeparator()
    a4 = QAction("Quit", menu)
    a4.triggered.connect(app.quit)
    menu.addAction(a4)
    tray.setContextMenu(menu)
    tray.activated.connect(lambda r: show_wheel_from_selection()
                           if r == QSystemTrayIcon.ActivationReason.Trigger else None)
    tray.show()

    apply_startup()
    QTimer.singleShot(100, hotkey_manager.restart)
    app.aboutToQuit.connect(hotkey_manager.stop)
    _toast("Format Wheel is running", [f"Hold {CONFIG['hotkey']} with files selected"], "info", 3000)
    sys.exit(app.exec())


if __name__ == "__main__":
    main()